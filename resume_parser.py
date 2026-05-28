import io
import json
import os

import docx
import fitz  # PyMuPDF
from langchain_groq import ChatGroq


def extract_text_from_upload(file_bytes: bytes, file_name: str) -> str:
    """Extract plain text from a PDF or DOCX file, putting hyperlinks inline."""
    ext = file_name.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            full_text = []
            for page in doc:
                words = page.get_text("words")
                links = page.get_links()
                
                uri_links = []
                for lnk in links:
                    if "uri" in lnk and "from" in lnk:
                        uri_links.append((fitz.Rect(lnk["from"]), lnk["uri"]))
                
                lines = {}
                for w in words:
                    x0, y0, x1, y1, word, block_no, line_no, word_no = w
                    key = (block_no, line_no)
                    if key not in lines:
                        lines[key] = []
                    
                    word_rect = fitz.Rect(x0, y0, x1, y1)
                    matched_uri = None
                    for link_rect, uri in uri_links:
                        if word_rect.intersects(link_rect) or word_rect in link_rect:
                            matched_uri = uri
                            break
                    lines[key].append((word, matched_uri))
                
                page_lines = []
                for key in sorted(lines.keys()):
                    words_in_line = lines[key]
                    line_str = ""
                    last_uri = None
                    for word, uri in words_in_line:
                        if uri and uri != last_uri:
                            if last_uri:
                                line_str += f" ({last_uri}) "
                            line_str += f" {word}"
                            last_uri = uri
                        elif not uri and last_uri:
                            line_str += f" ({last_uri}) {word}"
                            last_uri = None
                        else:
                            line_str += f" {word}"
                    if last_uri:
                        line_str += f" ({last_uri}) "
                    page_lines.append(line_str.strip())
                
                full_text.append("\n".join(page_lines))
            return "\n".join(full_text)
        except Exception as e:
            raise RuntimeError(f"Failed to read PDF: {e}")
    elif ext in ("docx", "doc"):
        try:
            document = docx.Document(io.BytesIO(file_bytes))
            paragraphs_text = []
            
            # Extract paragraphs (including inline hyperlinks)
            for para in document.paragraphs:
                p_text = ""
                for child in para._element:
                    tag = child.tag.split("}")[-1]
                    if tag == "r":
                        text_elems = child.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
                        p_text += "".join(t.text for t in text_elems if t.text)
                    elif tag == "hyperlink":
                        rId = child.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                        text_elems = child.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
                        link_text = "".join(t.text for t in text_elems if t.text)
                        
                        if rId and rId in para.part.rels:
                            rel = para.part.rels[rId]
                            url = getattr(rel, "target_ref", getattr(rel, "_target", ""))
                            p_text += f" {link_text} ({url}) "
                        else:
                            p_text += f" {link_text} "
                
                if p_text.strip():
                    paragraphs_text.append(p_text.strip())
            
            # Also get table text
            for table in document.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        paragraphs_text.append(" | ".join(row_text))

            return "\n".join(paragraphs_text)
        except Exception as e:
            raise RuntimeError(f"Failed to read DOCX: {e}")
    else:
        raise ValueError(f"Unsupported file type: .{ext}")


EXTRACTION_PROMPT = """
You are an expert resume parser. Extract ALL information from the resume text below
and return it as a single JSON object. Follow this schema exactly — use empty strings
or empty arrays when information is absent, never omit keys.

JSON Schema:
{{
  "Contact": {{
    "FirstName": "",
    "LastName": "",
    "Email": "",
    "Phone": "",
    "Location": "",
    "Profiles": [
      {{ "name": "LinkedIn", "link": "https://..." }}
    ]
  }},
  "Summary": "",
  "Objective": "",
  "Skills": "",
  "Experience": [
    {{
      "company": "",
      "location": "",
      "title": "",
      "dates": "",
      "bullets": ""
    }}
  ],
  "Education": [
    {{
      "school": "",
      "location": "",
      "degree": "",
      "dates": "",
      "gpa": "",
      "bullets": ""
    }}
  ],
  "Projects": [
    {{
      "title": "",
      "dates": "",
      "bullets": "",
      "link": ""
    }}
  ],
  "Certifications": "",
  "Interests": ""
}}

Rules:
- Skills: combine all skills into a single multi-line string (e.g. "Technical: Python, Java\\nLanguage: English")
- Certifications: one per line
- Bullets for experience/education/projects: one bullet per line, no leading dashes
- Project link: If a project title or description has an associated hyperlink or URL inline or nearby, extract that URL and set it in the project's "link" field.
- Profiles: only real URLs found in the resume
- Do NOT invent data. Leave fields empty if not found.

Resume text:
{resume_text}
"""


def parse_resume_with_groq(resume_text: str, api_key: str) -> dict:
    """Send resume text to Groq and return structured dict."""
    llm = ChatGroq(
        groq_api_key=api_key,
        model_name="llama-3.1-8b-instant",
        temperature=0.0,
    )
    llm_json = llm.bind(response_format={"type": "json_object"})

    prompt = EXTRACTION_PROMPT.format(resume_text=resume_text[:12000])  # token guard
    response = llm_json.invoke(prompt)
    raw = response.content if hasattr(response, "content") else str(response)

    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        # Try to extract JSON from the response
        start = raw.find("{")
        end = raw.rfind("}") + 1
        if start != -1 and end > start:
            return json.loads(raw[start:end])
        raise RuntimeError("Groq returned invalid JSON")


def fill_session_state(parsed: dict) -> None:
    """
    Populate st.session_state with extracted resume data so the editor
    form fields are pre-filled on the next rerun.
    """
    import streamlit as st

    def to_str(val) -> str:
        if isinstance(val, list):
            return "\n".join(str(v) for v in val if v)
        if val is None:
            return ""
        return str(val)

    contact = parsed.get("Contact", {})
    st.session_state.contact_first_name = to_str(contact.get("FirstName", ""))
    st.session_state.contact_last_name = to_str(contact.get("LastName", ""))
    st.session_state.contact_email = to_str(contact.get("Email", ""))
    st.session_state.contact_phone = to_str(contact.get("Phone", ""))
    st.session_state.contact_location = to_str(contact.get("Location", ""))

    profiles = contact.get("Profiles", [])
    if isinstance(profiles, list):
        st.session_state.prof_count = max(1, len(profiles))
        for i, prof in enumerate(profiles):
            if isinstance(prof, dict):
                st.session_state[f"prof_name_{i}"] = to_str(prof.get("name", ""))
                st.session_state[f"prof_link_{i}"] = to_str(prof.get("link", ""))
            else:
                st.session_state[f"prof_name_{i}"] = ""
                st.session_state[f"prof_link_{i}"] = to_str(prof)

    st.session_state.summary_input = to_str(parsed.get("Summary", ""))
    st.session_state.obj_input = to_str(parsed.get("Objective", ""))
    st.session_state.skills_input = to_str(parsed.get("Skills", ""))
    st.session_state.cert_input = to_str(parsed.get("Certifications", ""))
    st.session_state.int_input = to_str(parsed.get("Interests", ""))

    experiences = parsed.get("Experience", [])
    if isinstance(experiences, list):
        st.session_state.exp_count = max(1, len(experiences))
        for i, exp in enumerate(experiences):
            if isinstance(exp, dict):
                st.session_state[f"exp_org_{i}"] = to_str(exp.get("company", ""))
                st.session_state[f"exp_loc_{i}"] = to_str(exp.get("location", ""))
                st.session_state[f"exp_title_{i}"] = to_str(exp.get("title", ""))
                st.session_state[f"exp_dates_{i}"] = to_str(exp.get("dates", ""))
                st.session_state[f"exp_bull_{i}"] = to_str(exp.get("bullets", ""))

    educations = parsed.get("Education", [])
    if isinstance(educations, list):
        st.session_state.edu_count = max(1, len(educations))
        for i, edu in enumerate(educations):
            if isinstance(edu, dict):
                st.session_state[f"edu_school_{i}"] = to_str(edu.get("school", ""))
                st.session_state[f"edu_loc_{i}"] = to_str(edu.get("location", ""))
                st.session_state[f"edu_deg_{i}"] = to_str(edu.get("degree", ""))
                st.session_state[f"edu_dates_{i}"] = to_str(edu.get("dates", ""))
                st.session_state[f"edu_gpa_{i}"] = to_str(edu.get("gpa", ""))
                st.session_state[f"edu_bull_{i}"] = to_str(edu.get("bullets", ""))

    projects = parsed.get("Projects", [])
    if isinstance(projects, list):
        st.session_state.proj_count = max(1, len(projects))
        for i, proj in enumerate(projects):
            if isinstance(proj, dict):
                st.session_state[f"proj_title_{i}"] = to_str(proj.get("title", ""))
                st.session_state[f"proj_dates_{i}"] = to_str(proj.get("dates", ""))
                st.session_state[f"proj_desc_{i}"] = to_str(proj.get("bullets", ""))
                st.session_state[f"proj_link_{i}"] = to_str(proj.get("link", ""))

